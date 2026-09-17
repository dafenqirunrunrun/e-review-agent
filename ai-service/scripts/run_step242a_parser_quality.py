from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.document_ingestion.exceptions import DocumentIngestionError
from app.document_ingestion.router import DocumentParserRouter
from app.rag_quality.parser_metrics import aggregate_parser_results, evaluate_parsed_document


DEFAULT_FIXTURE_ROOT = ROOT / "artifacts" / "step242a" / "parser_fixtures"
DEFAULT_OUTPUT = ROOT / "artifacts" / "step242a" / "parser_quality_baseline.json"


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Run the Step 24.2A parser quality fixtures.")
    parser.add_argument("--docling-python", type=Path, required=True)
    parser.add_argument("--mineru-python", type=Path, required=True)
    parser.add_argument("--fixture-root", type=Path, default=DEFAULT_FIXTURE_ROOT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    fixture_root = args.fixture_root.resolve()
    fixture_root.mkdir(parents=True, exist_ok=True)
    os.environ["E_REVIEW_DOCLING_PYTHON"] = str(args.docling_python.resolve())
    os.environ["E_REVIEW_MINERU_PYTHON"] = str(args.mineru_python.resolve())
    os.environ["E_REVIEW_DOCUMENT_ARTIFACT_ROOT"] = str((fixture_root / "parsed_assets").resolve())

    fixtures = _create_fixtures(fixture_root, args.docling_python, args.mineru_python)
    router = DocumentParserRouter()
    results: list[dict[str, Any]] = []
    for path, expectation in fixtures:
        try:
            document = router.parse(
                path,
                document_id=f"step242a-{expectation['fixtureId']}",
                source_name=path.stem,
                source_type="rag_quality_fixture",
                language="zh",
            )
            result = evaluate_parsed_document(document, expectation)
            result["file"] = path.name
        except DocumentIngestionError as exc:
            result = _failed_result(path, expectation, exc.reason_code, str(exc))
        except Exception as exc:
            result = _failed_result(path, expectation, type(exc).__name__, str(exc))
        results.append(result)

    report = aggregate_parser_results(results)
    report.update(
        {
            "step": "24.2A",
            "fixtureRoot": str(fixture_root.relative_to(ROOT)).replace("\\", "/"),
            "formats": sorted({str(item[1]["expectedFormat"]) for item in fixtures}),
        }
    )
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "hardSafetyGate": report["hardSafetyGate"],
                "fixtureCount": report["fixtureCount"],
                "passedFixtureCount": report["passedFixtureCount"],
                "failures": [
                    {"fixtureId": item["fixtureId"], "checks": item.get("checks", {})}
                    for item in report["failures"]
                ],
                "output": str(output),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0 if report["hardSafetyGate"] == "PASS" else 1


def _failed_result(
    path: Path,
    expectation: dict[str, Any],
    reason_code: str,
    detail: str,
) -> dict[str, Any]:
    return {
        "fixtureId": expectation["fixtureId"],
        "file": path.name,
        "format": expectation["expectedFormat"],
        "parser": "none",
        "parserVersion": "none",
        "checks": {"parseSucceeded": False},
        "passed": False,
        "metrics": {},
        "nodeTypes": {},
        "warnings": [reason_code],
        "attempts": [],
        "error": detail[-1200:],
    }


def _create_fixtures(
    root: Path,
    docling_python: Path,
    mineru_python: Path,
) -> list[tuple[Path, dict[str, Any]]]:
    fixtures = [
        _write_html(root),
        _write_markdown(root),
        _write_text(root),
        _write_csv(root),
        _write_xlsx(root),
    ]
    fixtures.extend(_write_office(root, docling_python))
    fixtures.extend(_write_pdfs(root, mineru_python))
    return fixtures


def _write_html(root: Path) -> tuple[Path, dict[str, Any]]:
    path = root / "policy.html"
    path.write_text(
        "<h1>评论治理规范</h1><h2>有偿评价</h2>"
        "<p>商家不得通过返现诱导消费者提交五星评价。</p>"
        "<ol><li>保留交易证据</li><li>进入人工复核</li></ol>"
        "<table><tr><th>风险</th><th>处置</th></tr>"
        "<tr><td>评分操纵</td><td>人工复核</td></tr></table>",
        encoding="utf-8",
    )
    return path, _expect(
        "html-structure",
        "html",
        anchors=["返现", "人工复核"],
        ordered=["评论治理规范", "有偿评价", "五星评价"],
        node_types=["title", "section", "list_item", "table", "table_row"],
        section_paths=["评论治理规范 / 有偿评价"],
        table_cells=["评分操纵", "人工复核"],
        route_reasons=["TEXT_STRUCTURE_NATIVE"],
    )


def _write_markdown(root: Path) -> tuple[Path, dict[str, Any]]:
    path = root / "policy.md"
    path.write_text(
        "# 评论治理政策\n\n## 差评压制\n\n不得通过威胁方式要求消费者删除差评。\n\n"
        "1. 保留沟通记录\n2. 必要时人工复核\n\n"
        "| 风险 | 处置 |\n| --- | --- |\n| 删除差评 | 人工复核 |\n",
        encoding="utf-8",
    )
    return path, _expect(
        "markdown-structure",
        "markdown",
        anchors=["删除差评", "沟通记录"],
        ordered=["评论治理政策", "差评压制", "威胁"],
        node_types=["title", "section", "list_item", "table", "table_row"],
        section_paths=["评论治理政策 / 差评压制"],
        table_cells=["删除差评", "人工复核"],
        route_reasons=["TEXT_STRUCTURE_NATIVE"],
    )


def _write_text(root: Path) -> tuple[Path, dict[str, Any]]:
    path = root / "policy.txt"
    path.write_text("普通评价无需进入政策证据检索。\n\n缺少风险信号时保持低成本处理。\n", encoding="utf-8")
    return path, _expect(
        "text-basic",
        "text",
        anchors=["普通评价", "低成本处理"],
        ordered=["普通评价", "风险信号"],
        node_types=["paragraph"],
        route_reasons=["TEXT_STRUCTURE_NATIVE"],
    )


def _write_csv(root: Path) -> tuple[Path, dict[str, Any]]:
    path = root / "policy.csv"
    path.write_text("风险,阈值,处置\n虚假评价,高,人工复核\n评分操纵,中,保留证据\n", encoding="utf-8")
    return path, _expect(
        "csv-table",
        "csv",
        anchors=["虚假评价", "保留证据"],
        ordered=["风险", "虚假评价", "评分操纵"],
        node_types=["table", "table_row"],
        table_cells=["评分操纵", "人工复核"],
        route_reasons=["TEXT_STRUCTURE_NATIVE"],
    )


def _write_xlsx(root: Path) -> tuple[Path, dict[str, Any]]:
    import openpyxl

    path = root / "policy.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "风险规则"
    sheet.append(["风险", "阈值", "处置"])
    sheet.append(["虚假评价", 0.8, "人工复核"])
    sheet.append(["评分操纵", "=B2-0.1", "保留证据"])
    sheet.merge_cells("A5:C5")
    sheet["A5"] = "政策依据必须可追溯"
    workbook.save(path)
    workbook.close()
    return path, _expect(
        "xlsx-cell-graph",
        "xlsx",
        anchors=["虚假评价", "政策依据必须可追溯"],
        ordered=["风险规则", "虚假评价", "评分操纵"],
        node_types=["sheet", "table", "table_row"],
        section_paths=["风险规则"],
        table_cells=["评分操纵", "=B2-0.1"],
        route_reasons=["SPREADSHEET_CELL_GRAPH_REQUIRED"],
    )


def _write_office(root: Path, python: Path) -> list[tuple[Path, dict[str, Any]]]:
    docx = root / "policy.docx"
    pptx = root / "policy.pptx"
    code = r'''
from docx import Document
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
import sys

docx_path, pptx_path = map(Path, sys.argv[1:3])
document = Document()
document.add_heading('评论治理政策', 0)
document.add_heading('评价真实性', 1)
document.add_paragraph('好评返现可能影响评价真实性，应保留证据并进行审核。')
table = document.add_table(rows=2, cols=2)
table.cell(0, 0).text = '风险'
table.cell(0, 1).text = '建议'
table.cell(1, 0).text = '有偿评价'
table.cell(1, 1).text = '人工复核'
document.save(docx_path)

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[5])
slide.shapes.title.text = '评论治理演示'
box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8), Inches(0.8))
box.text_frame.text = '五星截图返现属于评分操纵风险'
grid = slide.shapes.add_table(2, 2, Inches(0.8), Inches(2.5), Inches(7), Inches(1.5)).table
grid.cell(0, 0).text = '风险'
grid.cell(0, 1).text = '处置'
grid.cell(1, 0).text = '评分操纵'
grid.cell(1, 1).text = '人工复核'
presentation.save(pptx_path)
'''
    _run(python, code, docx, pptx)
    return [
        (
            docx,
            _expect(
                "docx-layout",
                "docx",
                anchors=["好评返现", "人工复核"],
                ordered=["评论治理政策", "评价真实性", "好评返现"],
                node_types=["title", "section", "table"],
                section_paths=["评论治理政策", "评价真实性"],
                table_cells=["有偿评价", "人工复核"],
                route_reasons=["OFFICE_LAYOUT_REQUIRED"],
            ),
        ),
        (
            pptx,
            _expect(
                "pptx-slide-table",
                "pptx",
                anchors=["五星截图返现", "人工复核"],
                ordered=["评论治理演示", "五星截图返现", "评分操纵"],
                node_types=["title", "table"],
                table_cells=["评分操纵", "人工复核"],
                route_reasons=["OFFICE_LAYOUT_REQUIRED"],
            ),
        ),
    ]


def _write_pdfs(root: Path, python: Path) -> list[tuple[Path, dict[str, Any]]]:
    digital = root / "digital-policy.pdf"
    scanned = root / "scanned-policy.pdf"
    mixed = root / "mixed-policy.pdf"
    code = r'''
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
import sys

digital, scanned, mixed = map(Path, sys.argv[1:4])

c = canvas.Canvas(str(digital), pagesize=A4)
c.setFont('Helvetica-Bold', 18); c.drawString(60, 780, 'REVIEW GOVERNANCE POLICY')
c.setFont('Helvetica-Bold', 13); c.drawString(60, 740, 'PAID REVIEW')
c.setFont('Helvetica', 11); c.drawString(60, 710, 'Cashback for a five-star screenshot requires human review.')
c.rect(60, 600, 470, 70); c.line(280, 600, 280, 670); c.line(60, 635, 530, 635)
c.drawString(80, 650, 'RISK'); c.drawString(300, 650, 'ACTION')
c.drawString(80, 615, 'RATING MANIPULATION'); c.drawString(300, 615, 'HUMAN REVIEW')
c.showPage(); c.save()

image_path = scanned.with_suffix('.png')
image = Image.new('RGB', (1240, 1754), 'white')
draw = ImageDraw.Draw(image)
font = ImageFont.load_default(size=30); small = ImageFont.load_default(size=24)
draw.text((90, 90), 'REVIEW GOVERNANCE POLICY', fill='black', font=font)
draw.text((90, 180), 'PAID REVIEW AND RATING MANIPULATION', fill='black', font=small)
draw.text((90, 250), 'Cashback for a five-star screenshot requires human review.', fill='black', font=small)
draw.rectangle((90, 360, 1100, 620), outline='black', width=4)
draw.line((600, 360, 600, 620), fill='black', width=4); draw.line((90, 480, 1100, 480), fill='black', width=4)
draw.text((130, 400), 'RISK', fill='black', font=small); draw.text((650, 400), 'ACTION', fill='black', font=small)
draw.text((130, 530), 'FAKE REVIEW', fill='black', font=small); draw.text((650, 530), 'HUMAN REVIEW', fill='black', font=small)
image.save(image_path)
c = canvas.Canvas(str(scanned), pagesize=A4)
c.drawImage(str(image_path), 0, 0, width=A4[0], height=A4[1]); c.showPage(); c.save()

figure_path = mixed.with_suffix('.png')
figure = Image.new('RGB', (640, 280), 'white'); figure_draw = ImageDraw.Draw(figure)
figure_draw.rectangle((5, 5, 635, 275), outline='black', width=4)
figure_draw.text((35, 40), 'SCREENSHOT REWARD EVIDENCE', fill='black', font=small)
figure_draw.text((35, 120), 'RATING MANIPULATION', fill='black', font=small)
figure.save(figure_path)
c = canvas.Canvas(str(mixed), pagesize=A4)
c.setFont('Helvetica-Bold', 18); c.drawString(60, 780, 'MIXED POLICY EVIDENCE')
c.setFont('Helvetica', 11); c.drawString(60, 750, 'The policy image below is knowledge-bearing evidence.')
c.drawImage(str(figure_path), 60, 430, width=470, height=205)
c.showPage(); c.save()
'''
    _run(python, code, digital, scanned, mixed)
    return [
        (
            digital,
            _expect(
                "pdf-digital-table",
                "pdf",
                anchors=["PAID REVIEW", "HUMAN REVIEW"],
                ordered=["REVIEW GOVERNANCE POLICY", "PAID REVIEW", "Cashback"],
                node_types=["title", "table"],
                table_cells=["RATING MANIPULATION", "HUMAN REVIEW"],
                route_reasons=["PDF_TEXT_LAYOUT"],
                pages=1,
            ),
        ),
        (
            scanned,
            _expect(
                "pdf-scanned-ocr",
                "pdf",
                anchors=["PAID REVIEW", "HUMAN REVIEW"],
                ordered=["REVIEW GOVERNANCE POLICY", "PAID REVIEW", "Cashback"],
                node_types=["title", "section", "table"],
                route_reasons=["PDF_LAYOUT_OCR_COMPLEX"],
                pages=1,
            ),
        ),
        (
            mixed,
            _expect(
                "pdf-mixed-image",
                "pdf",
                anchors=["MIXED POLICY EVIDENCE", "knowledge-bearing evidence"],
                ordered=["MIXED POLICY EVIDENCE", "knowledge-bearing evidence"],
                node_types=["title", "figure"],
                route_reasons=["PDF_TEXT_LAYOUT"],
                pages=1,
                assets=1,
                asset_association=True,
            ),
        ),
    ]


def _expect(
    fixture_id: str,
    expected_format: str,
    *,
    anchors: list[str],
    ordered: list[str],
    node_types: list[str],
    section_paths: list[str] | None = None,
    table_cells: list[str] | None = None,
    route_reasons: list[str] | None = None,
    pages: int = 0,
    assets: int = 0,
    asset_association: bool = False,
) -> dict[str, Any]:
    return {
        "fixtureId": fixture_id,
        "expectedFormat": expected_format,
        "requiredAnchors": anchors,
        "orderedAnchors": ordered,
        "requiredNodeTypes": node_types,
        "requiredSectionPaths": section_paths or [],
        "requiredTableCells": table_cells or [],
        "requiredRouteReasons": route_reasons or [],
        "minimumNodeCount": 1,
        "minimumReferencedPageCount": pages,
        "minimumAssetCount": assets,
        "requireAssetAssociation": asset_association,
    }


def _run(python: Path, code: str, *paths: Path) -> None:
    completed = subprocess.run(
        [str(python.resolve()), "-c", code, *(str(path) for path in paths)],
        capture_output=True,
        text=True,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError((completed.stderr or completed.stdout or "FIXTURE_GENERATION_FAILED").strip())


if __name__ == "__main__":
    raise SystemExit(main())
