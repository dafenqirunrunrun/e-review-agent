from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.document_ingestion.router import DocumentParserRouter


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Step 24.1 document parser smoke verification")
    parser.add_argument("--docling-python", type=Path, required=True)
    parser.add_argument("--mineru-python", type=Path)
    args = parser.parse_args()

    os.environ["E_REVIEW_DOCLING_PYTHON"] = str(args.docling_python.resolve())
    if args.mineru_python:
        os.environ["E_REVIEW_MINERU_PYTHON"] = str(args.mineru_python.resolve())

    with tempfile.TemporaryDirectory(prefix="e-review-step241-") as temp_name:
        root = Path(temp_name)
        samples = _create_samples(root, args.docling_python, args.mineru_python)
        router = DocumentParserRouter()
        results: list[dict[str, object]] = []
        for sample in samples:
            document = router.parse(
                sample,
                document_id=f"step241-{sample.suffix.lstrip('.')}",
                source_name=sample.stem,
                source_type="verification_fixture",
                language="zh",
            )
            results.append(
                {
                    "file": sample.name,
                    "format": document.format,
                    "parser": document.parser,
                    "parserVersion": document.parserVersion,
                    "nodeCount": len(document.nodes),
                    "nodeTypes": dict(Counter(node.type for node in document.nodes)),
                    "sectionPaths": sorted({" / ".join(node.sectionPath) for node in document.nodes if node.sectionPath}),
                    "attempts": [attempt.model_dump(mode="json") for attempt in document.attempts],
                    "contentHashPresent": len(document.contentHash) == 64,
                    "markdownPresent": bool(document.markdown.strip()),
                    "semanticContentPresent": _semantic_content_present(sample.name, document.markdown),
                    "routeReasons": document.metadata["parserRoute"]["reasonCodes"],
                }
            )

    structure_valid = {
        "review-policy.html": {"title", "section", "list_item", "table", "table_row"},
        "review-policy.xlsx": {"sheet", "table", "table_row"},
        "review-policy.docx": {"title", "section", "table"},
    }
    if args.mineru_python:
        structure_valid["scanned-review-policy.pdf"] = {"paragraph"}
    passed = all(
        item["contentHashPresent"]
        and item["markdownPresent"]
        and int(item["nodeCount"]) > 0
        and structure_valid[str(item["file"])].issubset(set(item["nodeTypes"]))
        and item["semanticContentPresent"]
        and (
            str(item["file"]) != "scanned-review-policy.pdf"
            or item["routeReasons"] == ["PDF_LAYOUT_OCR_COMPLEX"]
        )
        for item in results
    )
    print(json.dumps({"step": "24.1", "passed": passed, "results": results}, ensure_ascii=False, indent=2))
    return 0 if passed else 1


def _create_samples(root: Path, docling_python: Path, mineru_python: Path | None) -> list[Path]:
    html_path = root / "review-policy.html"
    html_path.write_text(
        """
<!doctype html>
<html><body>
<h1>评论治理规范</h1>
<h2>有偿评价</h2>
<p>商家不得通过返现诱导消费者提交五星评价。</p>
<ul><li>保留交易证据</li><li>必要时进入人工复核</li></ul>
<table><tr><th>风险</th><th>处理</th></tr><tr><td>评分操纵</td><td>人工复核</td></tr></table>
</body></html>
""".strip(),
        encoding="utf-8",
    )

    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("OPENPYXL_REQUIRED_FOR_STEP241_SMOKE") from exc
    xlsx_path = root / "review-policy.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "风险规则"
    sheet.append(["风险", "阈值", "处置"])
    sheet.append(["虚假评价", 0.8, "人工复核"])
    sheet.append(["评分操纵", "=B2-0.1", "限制自动处置"])
    sheet.merge_cells("A5:C5")
    sheet["A5"] = "政策依据必须可追溯"
    workbook.save(xlsx_path)
    workbook.close()

    docx_path = root / "review-policy.docx"
    _write_sample_docx(docx_path, docling_python)
    samples = [html_path, xlsx_path, docx_path]
    if mineru_python:
        scanned_pdf = root / "scanned-review-policy.pdf"
        _write_scanned_pdf(scanned_pdf, mineru_python)
        samples.append(scanned_pdf)
    return samples


def _write_sample_docx(path: Path, docling_python: Path) -> None:
    code = """
from docx import Document
from pathlib import Path
import sys
document = Document()
document.add_heading('评论治理政策', 0)
document.add_heading('评价真实性', 1)
document.add_paragraph('好评返现可能影响评价真实性，应保留证据并进行审核。')
table = document.add_table(rows=2, cols=2)
table.cell(0, 0).text = '风险'
table.cell(0, 1).text = '建议'
table.cell(1, 0).text = '有偿评价'
table.cell(1, 1).text = '人工复核'
document.save(Path(sys.argv[1]))
"""
    completed = subprocess.run(
        [str(docling_python.resolve()), "-c", code, str(path)],
        capture_output=True,
        text=True,
        check=False,
    )
    if completed.returncode != 0 or not path.is_file():
        raise RuntimeError(f"DOCX_FIXTURE_GENERATION_FAILED: {completed.stderr.strip()}")


def _write_scanned_pdf(path: Path, mineru_python: Path) -> None:
    code = """
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from pathlib import Path
import sys

target = Path(sys.argv[1])
image_path = target.with_suffix('.png')
image = Image.new('RGB', (1240, 1754), 'white')
draw = ImageDraw.Draw(image)
font = ImageFont.load_default(size=30)
small = ImageFont.load_default(size=24)
draw.text((90, 90), 'REVIEW GOVERNANCE POLICY', fill='black', font=font)
draw.text((90, 180), 'PAID REVIEW AND RATING MANIPULATION', fill='black', font=small)
draw.text((90, 250), 'Cashback for a five-star screenshot requires manual review.', fill='black', font=small)
draw.rectangle((90, 360, 1100, 620), outline='black', width=4)
draw.line((600, 360, 600, 620), fill='black', width=4)
draw.line((90, 480, 1100, 480), fill='black', width=4)
draw.text((130, 400), 'RISK', fill='black', font=small)
draw.text((650, 400), 'ACTION', fill='black', font=small)
draw.text((130, 530), 'FAKE REVIEW', fill='black', font=small)
draw.text((650, 530), 'HUMAN REVIEW', fill='black', font=small)
image.save(image_path)
c = canvas.Canvas(str(target), pagesize=A4)
c.drawImage(str(image_path), 0, 0, width=A4[0], height=A4[1])
c.showPage()
c.save()
"""
    completed = subprocess.run(
        [str(mineru_python.resolve()), "-c", code, str(path)],
        capture_output=True,
        text=True,
        check=False,
    )
    if completed.returncode != 0 or not path.is_file():
        raise RuntimeError(f"SCANNED_PDF_FIXTURE_GENERATION_FAILED: {completed.stderr.strip()}")


def _semantic_content_present(file_name: str, markdown: str) -> bool:
    lowered = markdown.lower()
    expected_terms = {
        "review-policy.html": ("返现", "人工复核"),
        "review-policy.xlsx": ("虚假评价", "政策依据"),
        "review-policy.docx": ("好评返现", "人工复核"),
        "scanned-review-policy.pdf": ("paid review", "human review"),
    }
    return all(term in lowered for term in expected_terms[file_name])


if __name__ == "__main__":
    raise SystemExit(main())
