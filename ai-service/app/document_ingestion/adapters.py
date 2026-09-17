from __future__ import annotations

import csv
import html
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from abc import ABC, abstractmethod
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

from app.document_ingestion.exceptions import DocumentParseError, ParserUnavailableError
from app.document_ingestion.models import (
    BoundingBox,
    DocumentAsset,
    DocumentFormat,
    DocumentNode,
    DocumentProbe,
    NormalizedDocument,
    ParserKind,
)
from app.rag.document_contract import stable_hash


TEXT_FORMATS = {"html", "markdown", "text", "csv"}


class DocumentParserAdapter(ABC):
    kind: ParserKind

    @abstractmethod
    def available(self) -> bool:
        raise NotImplementedError

    @abstractmethod
    def parse(
        self,
        path: Path,
        *,
        probe: DocumentProbe,
        document_id: str,
        source_name: str,
        source_type: str,
        source_uri: str,
        language: str,
    ) -> NormalizedDocument:
        raise NotImplementedError


class LightweightDocumentAdapter(DocumentParserAdapter):
    kind: ParserKind = "lightweight"

    def available(self) -> bool:
        return True

    def parse(
        self,
        path: Path,
        *,
        probe: DocumentProbe,
        document_id: str,
        source_name: str,
        source_type: str,
        source_uri: str,
        language: str,
    ) -> NormalizedDocument:
        if probe.format not in TEXT_FORMATS:
            raise DocumentParseError(
                "LIGHTWEIGHT_BINARY_UNSUPPORTED",
                f"Lightweight parsing is not allowed for {probe.format} files",
            )
        try:
            text = path.read_text(encoding="utf-8-sig")
        except (UnicodeError, OSError) as exc:
            raise DocumentParseError("LIGHTWEIGHT_TEXT_READ_FAILED", str(exc)) from exc

        if probe.format == "html":
            markdown, nodes = _parse_html(text, document_id)
        elif probe.format == "csv":
            markdown, nodes = _parse_csv(path, document_id)
        else:
            markdown = _normalize_markdown(text)
            nodes = _nodes_from_markdown(markdown, document_id)

        return _document(
            path=path,
            probe=probe,
            document_id=document_id,
            source_name=source_name,
            source_type=source_type,
            source_uri=source_uri,
            language=language,
            parser=self.kind,
            parser_version="stdlib-v1",
            markdown=markdown,
            nodes=nodes,
            metadata={"structureSource": "lightweight"},
        )


class SpreadsheetDocumentAdapter(DocumentParserAdapter):
    kind: ParserKind = "spreadsheet_native"

    def available(self) -> bool:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            return False
        return True

    def parse(
        self,
        path: Path,
        *,
        probe: DocumentProbe,
        document_id: str,
        source_name: str,
        source_type: str,
        source_uri: str,
        language: str,
    ) -> NormalizedDocument:
        if probe.format != "xlsx":
            raise DocumentParseError("SPREADSHEET_FORMAT_UNSUPPORTED", probe.format)
        try:
            import openpyxl
        except ImportError as exc:
            raise ParserUnavailableError("OPENPYXL_UNAVAILABLE") from exc

        try:
            workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
        except Exception as exc:
            raise DocumentParseError("XLSX_PARSE_FAILED", str(exc)) from exc

        nodes: list[DocumentNode] = []
        markdown_parts: list[str] = []
        order = 0
        sheet_count = len(workbook.worksheets)
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                sheet_id = _node_id(document_id, "sheet", sheet_index, sheet.title)
                table_id = _node_id(document_id, "table", sheet_index, sheet.title)
                sheet_order = order
                order += 1
                table_order = order
                order += 1
                row_nodes: list[DocumentNode] = []
                markdown_rows: list[list[str]] = []
                max_columns = 0

                for row_index, row in enumerate(sheet.iter_rows(), start=1):
                    populated: list[dict[str, Any]] = []
                    display_values: list[str] = []
                    for cell in row:
                        value = cell.value
                        display = "" if value is None else str(value)
                        display_values.append(display)
                        if value is not None:
                            populated.append(
                                {
                                    "coordinate": cell.coordinate,
                                    "value": display,
                                    "dataType": cell.data_type,
                                    "isFormula": cell.data_type == "f" or display.startswith("="),
                                }
                            )
                    while display_values and not display_values[-1]:
                        display_values.pop()
                    if not populated:
                        continue
                    max_columns = max(max_columns, len(display_values))
                    markdown_rows.append(display_values)
                    row_id = _node_id(document_id, "row", sheet_index, row_index, json.dumps(populated, ensure_ascii=False))
                    row_nodes.append(
                        DocumentNode(
                            nodeId=row_id,
                            type="table_row",
                            order=order,
                            text=" | ".join(display_values),
                            parentId=table_id,
                            sectionPath=[sheet.title],
                            sourceRef=f"{sheet.title}!{row[0].coordinate}:{row[-1].coordinate}",
                            metadata={"rowNumber": row_index, "cells": populated},
                        )
                    )
                    order += 1

                table_children = [item.nodeId for item in row_nodes]
                sheet_children = [table_id] if row_nodes else []
                nodes.append(
                    DocumentNode(
                        nodeId=sheet_id,
                        type="sheet",
                        order=sheet_order,
                        text=sheet.title,
                        children=sheet_children,
                        sectionPath=[sheet.title],
                        sourceRef=sheet.title,
                        metadata={
                            "sheetIndex": sheet_index,
                            "mergedRanges": [str(item) for item in sheet.merged_cells.ranges],
                            "maxRow": sheet.max_row,
                            "maxColumn": sheet.max_column,
                        },
                    )
                )
                if row_nodes:
                    nodes.append(
                        DocumentNode(
                            nodeId=table_id,
                            type="table",
                            order=table_order,
                            parentId=sheet_id,
                            children=table_children,
                            sectionPath=[sheet.title],
                            sourceRef=sheet.title,
                            metadata={"rowCount": len(row_nodes), "columnCount": max_columns},
                        )
                    )
                    nodes.extend(row_nodes)

                markdown_parts.append(f"# {sheet.title}")
                if markdown_rows:
                    width = max(len(row) for row in markdown_rows)
                    padded = [row + [""] * (width - len(row)) for row in markdown_rows]
                    markdown_parts.append(_markdown_table(padded))
        finally:
            workbook.close()

        markdown = "\n\n".join(part for part in markdown_parts if part).strip()
        return _document(
            path=path,
            probe=probe,
            document_id=document_id,
            source_name=source_name,
            source_type=source_type,
            source_uri=source_uri,
            language=language,
            parser=self.kind,
            parser_version=getattr(openpyxl, "__version__", "unknown"),
            markdown=markdown,
            nodes=nodes,
            metadata={"sheetCount": sheet_count, "structureSource": "openpyxl"},
        )


class PdfTextDocumentAdapter(DocumentParserAdapter):
    """Fast path for born-digital PDFs whose structure does not require OCR."""

    kind: ParserKind = "pdf_text"

    def available(self) -> bool:
        try:
            import pypdf  # noqa: F401
        except ImportError:
            return False
        return True

    def parse(
        self,
        path: Path,
        *,
        probe: DocumentProbe,
        document_id: str,
        source_name: str,
        source_type: str,
        source_uri: str,
        language: str,
    ) -> NormalizedDocument:
        if probe.format != "pdf" or probe.requiresOcr or probe.hasTextLayer is False:
            raise DocumentParseError("PDF_TEXT_FAST_PATH_UNSUPPORTED", probe.format)
        try:
            import pypdf
            from pypdf import PdfReader
        except ImportError as exc:
            raise ParserUnavailableError("PYPDF_UNAVAILABLE") from exc

        try:
            reader = PdfReader(path, strict=False)
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise DocumentParseError("PYPDF_TEXT_PARSE_FAILED", str(exc)) from exc
        if sum(len(re.sub(r"\s+", "", page)) for page in pages) < 20:
            raise DocumentParseError("PYPDF_TEXT_LAYER_EMPTY")

        markdown, removed_lines = _pdf_pages_to_markdown(pages, source_name)
        nodes = _nodes_from_markdown(markdown, document_id)
        return _document(
            path=path,
            probe=probe,
            document_id=document_id,
            source_name=source_name,
            source_type=source_type,
            source_uri=source_uri,
            language=language,
            parser=self.kind,
            parser_version=getattr(pypdf, "__version__", "unknown"),
            markdown=markdown,
            nodes=nodes,
            metadata={
                "pageCount": len(pages),
                "removedRepeatedLayoutLines": removed_lines,
                "structureSource": "pypdf_text",
            },
        )


class ExternalWorkerDocumentAdapter(DocumentParserAdapter):
    """Runs a heavyweight parser in an isolated Python environment.

    The worker returns a small parser-neutral JSON envelope. This keeps Docling
    and MinerU dependencies out of the always-on review service environment.
    """

    def __init__(
        self,
        kind: ParserKind,
        *,
        python_executable: str | Path | None,
        worker_script: str | Path,
        artifact_root: str | Path | None = None,
        timeout_seconds: int = 600,
    ):
        if kind not in {"docling", "mineru"}:
            raise ValueError("EXTERNAL_PARSER_KIND_INVALID")
        self.kind = kind
        self.python_executable = Path(python_executable).expanduser() if python_executable else None
        self.worker_script = Path(worker_script).expanduser()
        self.artifact_root = Path(artifact_root).expanduser() if artifact_root else None
        self.timeout_seconds = timeout_seconds

    def available(self) -> bool:
        return bool(
            self.python_executable
            and self.python_executable.is_file()
            and self.worker_script.is_file()
        )

    def parse(
        self,
        path: Path,
        *,
        probe: DocumentProbe,
        document_id: str,
        source_name: str,
        source_type: str,
        source_uri: str,
        language: str,
    ) -> NormalizedDocument:
        if not self.available():
            raise ParserUnavailableError(f"{self.kind.upper()}_WORKER_UNAVAILABLE")

        runtime_root = _external_worker_runtime_root(self.worker_script)
        with tempfile.TemporaryDirectory(prefix=f"e-review-{self.kind}-", dir=runtime_root) as temp_name:
            temp_dir = Path(temp_name)
            output_path = temp_dir / "result.json"
            bundle_dir = temp_dir / "bundle"
            command = [
                str(self.python_executable),
                str(self.worker_script),
                "--parser",
                self.kind,
                "--input",
                str(path),
                "--output",
                str(output_path),
                "--bundle-dir",
                str(bundle_dir),
            ]
            started = time.perf_counter()
            try:
                completed = subprocess.run(
                    command,
                    capture_output=True,
                    text=True,
                    timeout=self.timeout_seconds,
                    check=False,
                    env=_external_worker_environment(self.kind, runtime_root),
                )
            except (OSError, subprocess.TimeoutExpired) as exc:
                raise DocumentParseError(f"{self.kind.upper()}_WORKER_EXECUTION_FAILED", str(exc)) from exc
            duration_ms = round((time.perf_counter() - started) * 1000, 3)
            if completed.returncode != 0 or not output_path.is_file():
                detail = (completed.stderr or completed.stdout or "worker produced no result").strip()[-1200:]
                raise DocumentParseError(f"{self.kind.upper()}_PARSE_FAILED", detail)

            try:
                payload = json.loads(output_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                raise DocumentParseError(f"{self.kind.upper()}_RESULT_INVALID", str(exc)) from exc

            markdown = _normalize_markdown(str(payload.get("markdown", "")))
            nodes = _nodes_from_worker(payload.get("elements", []), document_id, markdown)
            assets = _assets_from_worker(payload.get("assets", []), document_id, bundle_dir, self.artifact_root)
            known_assets = {asset.assetId for asset in assets}
            nodes = [
                node.model_copy(update={"assetIds": [item for item in node.assetIds if item in known_assets]})
                for node in nodes
            ]
            metadata = dict(payload.get("metadata") or {})
            metadata.update({"workerDurationMs": duration_ms, "structureSource": self.kind})
            return _document(
                path=path,
                probe=probe,
                document_id=document_id,
                source_name=source_name,
                source_type=source_type,
                source_uri=source_uri,
                language=language,
                parser=self.kind,
                parser_version=str(payload.get("parserVersion", "unknown")),
                markdown=markdown,
                nodes=nodes,
                assets=assets,
                warnings=[str(item) for item in payload.get("warnings", [])],
                metadata=metadata,
            )


class _StructuralHtmlParser(HTMLParser):
    _block_tags = {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "td", "th", "caption"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._tag: str | None = None
        self._parts: list[str] = []
        self._skip_depth = 0
        self._skipped_tags: list[str] = []
        self._in_row = False
        self._row_cells: list[str] = []
        self.items: list[tuple[str, str]] = []
        self.scoped_items: list[tuple[str, str]] = []
        self._content_scopes: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        attributes = {key.lower(): (value or "").lower() for key, value in attrs}
        classes = attributes.get("class", "")
        if tag in {"main", "article"}:
            self._content_scopes.append(tag)
        decorative_header = tag == "header" and (
            attributes.get("role") == "banner" or any(term in classes for term in ("site-header", "global-header"))
        )
        if tag in {"script", "style", "nav", "footer", "aside", "form", "noscript", "svg", "dialog"} or decorative_header:
            self._skip_depth += 1
            self._skipped_tags.append(tag)
            return
        if self._skip_depth:
            return
        if tag == "tr":
            self._flush()
            self._in_row = True
            self._row_cells = []
            return
        if tag in self._block_tags:
            self._flush()
            self._tag = tag
        elif tag == "br" and self._tag:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self._skipped_tags and tag == self._skipped_tags[-1]:
            self._skipped_tags.pop()
            self._skip_depth = max(0, self._skip_depth - 1)
            if self._content_scopes and self._content_scopes[-1] == tag:
                self._content_scopes.pop()
            return
        if not self._skip_depth and tag == "tr":
            self._flush()
            if self._row_cells:
                self.items.append(("tr", "\t".join(self._row_cells)))
            self._in_row = False
            self._row_cells = []
            return
        if not self._skip_depth and tag == self._tag:
            self._flush()
        if self._content_scopes and self._content_scopes[-1] == tag:
            self._content_scopes.pop()

    def handle_data(self, data: str) -> None:
        if not self._skip_depth and self._tag:
            self._parts.append(data)

    def close(self) -> None:
        super().close()
        self._flush()

    def _flush(self) -> None:
        if self._tag:
            value = re.sub(r"\s+", " ", "".join(self._parts)).strip()
            if value:
                decoded = html.unescape(value)
                if self._in_row and self._tag in {"td", "th"}:
                    self._row_cells.append(decoded)
                else:
                    self.items.append((self._tag, decoded))
                    if self._content_scopes:
                        self.scoped_items.append((self._tag, decoded))
        self._tag = None
        self._parts = []


def _parse_html(content: str, document_id: str) -> tuple[str, list[DocumentNode]]:
    parser = _StructuralHtmlParser()
    parser.feed(content)
    parser.close()
    scoped_chars = sum(len(text) for _, text in parser.scoped_items)
    items = parser.scoped_items if len(parser.scoped_items) >= 2 and scoped_chars >= 40 else parser.items
    markdown_lines: list[str] = []
    previous_tag = ""
    for tag, text in items:
        prefix = "\n" if tag == "tr" and previous_tag == "tr" else "\n\n"
        if tag.startswith("h") and tag[1:].isdigit():
            rendered = f"{'#' * int(tag[1:])} {text}"
        elif tag == "li":
            rendered = f"- {text}"
        elif tag == "tr":
            rendered = "| " + " | ".join(text.split("\t")) + " |"
        elif tag == "caption":
            rendered = f"*{text}*"
        else:
            rendered = text
        markdown_lines.append((prefix if markdown_lines else "") + rendered)
        previous_tag = tag
    markdown = _normalize_markdown("".join(markdown_lines))
    return markdown, _nodes_from_markdown(markdown, document_id)


def _pdf_pages_to_markdown(pages: list[str], source_name: str) -> tuple[str, int]:
    cleaned_pages = [[re.sub(r"\s+", " ", line).strip() for line in page.splitlines()] for page in pages]
    edge_lines: Counter[str] = Counter()
    for lines in cleaned_pages:
        non_empty = [line for line in lines if line]
        edge_lines.update(set([*non_empty[:3], *non_empty[-3:]]))
    repeat_threshold = max(3, int(len(pages) * 0.4))
    repeated = {line for line, count in edge_lines.items() if count >= repeat_threshold and len(line) <= 160}
    removed = 0
    parts = [f"# {source_name}"]
    for lines in cleaned_pages:
        page_lines = []
        for line in lines:
            if not line or PAGE_NUMBER_LINE_RE.fullmatch(line) or line in repeated:
                removed += int(bool(line))
                continue
            page_lines.append(line.replace("\ufffd", ""))
        rendered = _pdf_lines_to_markdown(page_lines)
        if rendered:
            parts.append(rendered)
    return _normalize_markdown("\n\n".join(parts)), removed


PAGE_NUMBER_LINE_RE = re.compile(r"^(?:page\s+)?\d+(?:\s+of\s+\d+)?$", re.IGNORECASE)
PDF_CLAUSE_LINE_RE = re.compile(r"^(?:§\s*\d|Article\s+\d|第\s*[一二三四五六七八九十百千万0-9]+\s*条)", re.IGNORECASE)
PDF_LIST_LINE_RE = re.compile(r"^(?:[-*•]|\(?[0-9A-Za-z]+\)|[0-9A-Za-z]+[.)])\s+")


def _pdf_lines_to_markdown(lines: list[str]) -> str:
    output: list[str] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            output.append(" ".join(paragraph).strip())
            paragraph.clear()

    for line in lines:
        heading = len(line) <= 120 and (
            (line.isupper() and any(char.isalpha() for char in line))
            or re.match(r"^(?:PART|SECTION)\s+[0-9A-Z]", line, re.IGNORECASE)
        )
        if heading:
            flush()
            output.append(f"## {line}")
        elif PDF_CLAUSE_LINE_RE.match(line) or PDF_LIST_LINE_RE.match(line):
            flush()
            paragraph.append(line)
        else:
            paragraph.append(line)
        if paragraph and line.endswith((".", "?", "!", ";", "。", "？", "！", "；")):
            flush()
    flush()
    return "\n\n".join(output)


def _parse_csv(path: Path, document_id: str) -> tuple[str, list[DocumentNode]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = [[str(cell) for cell in row] for row in csv.reader(handle)]
    except (UnicodeError, OSError, csv.Error) as exc:
        raise DocumentParseError("CSV_PARSE_FAILED", str(exc)) from exc
    table_id = _node_id(document_id, "csv-table", 0)
    row_nodes = [
        DocumentNode(
            nodeId=_node_id(document_id, "csv-row", index, json.dumps(row, ensure_ascii=False)),
            type="table_row",
            order=index + 1,
            text=" | ".join(row),
            parentId=table_id,
            sourceRef=f"row:{index + 1}",
            metadata={"cells": row, "rowNumber": index + 1},
        )
        for index, row in enumerate(rows)
        if any(cell.strip() for cell in row)
    ]
    table = DocumentNode(
        nodeId=table_id,
        type="table",
        order=0,
        children=[row.nodeId for row in row_nodes],
        sourceRef="csv:table",
        metadata={"rowCount": len(row_nodes), "columnCount": max((len(row) for row in rows), default=0)},
    )
    return _markdown_table(rows), [table, *row_nodes]


def _nodes_from_markdown(markdown: str, document_id: str) -> list[DocumentNode]:
    nodes: list[DocumentNode] = []
    heading_stack: list[str] = []
    paragraph: list[str] = []
    table_rows: list[list[str]] = []
    table_section_path: list[str] = []
    order = 0

    def append(node_type: str, text: str, *, metadata: dict[str, Any] | None = None) -> None:
        nonlocal order
        cleaned = text.strip()
        if not cleaned:
            return
        nodes.append(
            DocumentNode(
                nodeId=_node_id(document_id, node_type, order, cleaned),
                type=node_type,  # type: ignore[arg-type]
                order=order,
                text=cleaned,
                sectionPath=list(heading_stack),
                sourceRef=f"markdown:{order}",
                metadata=metadata or {},
            )
        )
        order += 1

    def flush_paragraph() -> None:
        nonlocal paragraph
        append("paragraph", "\n".join(paragraph))
        paragraph = []

    def flush_table() -> None:
        nonlocal table_rows, table_section_path, order
        if not table_rows:
            return
        table_text = _markdown_table(table_rows)
        table_id = _node_id(document_id, "table", order, table_text)
        table_order = order
        order += 1
        row_nodes: list[DocumentNode] = []
        for row_index, cells in enumerate(table_rows):
            row_nodes.append(
                DocumentNode(
                    nodeId=_node_id(document_id, "table-row", order, json.dumps(cells, ensure_ascii=False)),
                    type="table_row",
                    order=order,
                    text=" | ".join(cells),
                    parentId=table_id,
                    sectionPath=list(table_section_path),
                    sourceRef=f"markdown-table:{table_order}:row:{row_index}",
                    metadata={"cells": cells, "rowNumber": row_index + 1},
                )
            )
            order += 1
        nodes.append(
            DocumentNode(
                nodeId=table_id,
                type="table",
                order=table_order,
                text=table_text,
                children=[row.nodeId for row in row_nodes],
                sectionPath=list(table_section_path),
                sourceRef=f"markdown-table:{table_order}",
                metadata={"rowCount": len(row_nodes), "columnCount": max(len(row) for row in table_rows)},
            )
        )
        nodes.extend(row_nodes)
        table_rows = []
        table_section_path = []

    for raw in markdown.splitlines():
        line = raw.strip()
        if not line:
            flush_paragraph()
            flush_table()
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            flush_table()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_stack = heading_stack[: level - 1] + [title]
            append("title" if level == 1 else "section", title, metadata={"level": level})
        elif re.match(r"^(?:[-*+]\s+|\d+[.)]\s+)", line):
            flush_paragraph()
            flush_table()
            append("list_item", re.sub(r"^(?:[-*+]\s+|\d+[.)]\s+)", "", line))
        elif line.startswith("|") and line.endswith("|"):
            flush_paragraph()
            if not re.match(r"^\|\s*:?-+", line):
                cells = [cell.strip() for cell in line.strip("|").split("|")]
                if not table_rows:
                    table_section_path = list(heading_stack)
                table_rows.append(cells)
        else:
            flush_table()
            paragraph.append(line)
    flush_paragraph()
    flush_table()
    return nodes


def _nodes_from_worker(elements: Iterable[dict[str, Any]], document_id: str, markdown: str) -> list[DocumentNode]:
    nodes: list[DocumentNode] = []
    for order, raw in enumerate(elements):
        node_type = _map_node_type(str(raw.get("type", "paragraph")))
        bbox_payload = raw.get("bbox")
        bbox = None
        if isinstance(bbox_payload, dict):
            try:
                bbox = BoundingBox.model_validate(bbox_payload)
            except ValueError:
                bbox = None
        asset_ids = [
            _asset_id(document_id, str(item))
            for item in raw.get("assetRefs", [])
            if str(item).strip()
        ]
        nodes.append(
            DocumentNode(
                nodeId=_node_id(document_id, node_type, order, str(raw.get("sourceRef", "")), str(raw.get("text", ""))),
                type=node_type,
                order=order,
                text=str(raw.get("text", "")),
                sectionPath=[str(item) for item in raw.get("sectionPath", []) if str(item).strip()],
                pageNumber=_positive_int(raw.get("pageNumber")),
                bbox=bbox,
                assetIds=asset_ids,
                confidence=_confidence(raw.get("confidence")),
                sourceRef=str(raw.get("sourceRef", "")),
                metadata=dict(raw.get("metadata") or {}),
            )
        )
    return nodes or _nodes_from_markdown(markdown, document_id)


def _assets_from_worker(
    raw_assets: Iterable[dict[str, Any]],
    document_id: str,
    bundle_dir: Path,
    artifact_root: Path | None,
) -> list[DocumentAsset]:
    assets: list[DocumentAsset] = []
    for raw in raw_assets:
        source_ref = str(raw.get("sourceRef", raw.get("relativePath", "")))
        asset_id = _asset_id(document_id, source_ref)
        relative_path = str(raw.get("relativePath", ""))
        source_path = (bundle_dir / relative_path).resolve() if relative_path else None
        saved_relative = ""
        content = str(raw.get("ocrText", raw.get("caption", source_ref)))
        if source_path and source_path.is_file() and artifact_root:
            destination_dir = artifact_root / document_id / "assets"
            destination_dir.mkdir(parents=True, exist_ok=True)
            destination = destination_dir / source_path.name
            shutil.copy2(source_path, destination)
            saved_relative = destination.relative_to(artifact_root).as_posix()
            content_hash = stable_hash(destination.read_bytes())
        elif source_path and source_path.is_file():
            content_hash = stable_hash(source_path.read_bytes())
        else:
            content_hash = stable_hash(content)
        assets.append(
            DocumentAsset(
                assetId=asset_id,
                kind=_map_asset_kind(str(raw.get("kind", "image"))),
                mediaType=str(raw.get("mediaType", "application/octet-stream")),
                relativePath=saved_relative,
                pageNumber=_positive_int(raw.get("pageNumber")),
                caption=str(raw.get("caption", "")),
                ocrText=str(raw.get("ocrText", "")),
                description=str(raw.get("description", "")),
                knowledgeBearing=bool(raw.get("knowledgeBearing", True)),
                contentHash=content_hash,
                metadata=dict(raw.get("metadata") or {}),
            )
        )
    return assets


def _document(
    *,
    path: Path,
    probe: DocumentProbe,
    document_id: str,
    source_name: str,
    source_type: str,
    source_uri: str,
    language: str,
    parser: ParserKind,
    parser_version: str,
    markdown: str,
    nodes: list[DocumentNode],
    assets: list[DocumentAsset] | None = None,
    warnings: list[str] | None = None,
    metadata: dict[str, Any] | None = None,
) -> NormalizedDocument:
    resolved_assets = assets or []
    return NormalizedDocument(
        documentId=document_id,
        sourceName=source_name,
        sourceType=source_type,
        sourceUri=source_uri,
        format=probe.format,
        mimeType=probe.mimeType,
        language=language,
        parser=parser,
        parserVersion=parser_version,
        contentHash=stable_hash(path.read_bytes()),
        markdown=markdown,
        nodes=nodes,
        assets=resolved_assets,
        warnings=warnings or [],
        metadata={"normalizedContentHash": NormalizedDocument.content_digest(markdown=markdown, nodes=nodes, assets=resolved_assets), **(metadata or {})},
    )


def _normalize_markdown(content: str) -> str:
    text = content.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    if width == 0:
        return ""
    padded = [row + [""] * (width - len(row)) for row in rows]
    header = padded[0]
    body = padded[1:]
    lines = ["| " + " | ".join(_escape_table_cell(item) for item in header) + " |"]
    lines.append("| " + " | ".join("---" for _ in header) + " |")
    lines.extend("| " + " | ".join(_escape_table_cell(item) for item in row) + " |" for row in body)
    return "\n".join(lines)


def _escape_table_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ").strip()


def _node_id(document_id: str, *parts: object) -> str:
    return stable_hash(":".join([document_id, *(str(item) for item in parts)]))[:24]


def _asset_id(document_id: str, source_ref: str) -> str:
    return stable_hash(f"{document_id}:asset:{source_ref}")[:24]


def _positive_int(value: Any) -> int | None:
    try:
        number = int(value)
    except (TypeError, ValueError):
        return None
    return number if number >= 1 else None


def _confidence(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return min(1.0, max(0.0, number))


def _map_node_type(value: str) -> str:
    normalized = value.lower().replace("-", "_")
    mapping = {
        "heading": "section",
        "title": "title",
        "section": "section",
        "section_header": "section",
        "text": "paragraph",
        "paragraph": "paragraph",
        "list": "list_item",
        "list_item": "list_item",
        "table": "table",
        "table_row": "table_row",
        "picture": "figure",
        "image": "figure",
        "figure": "figure",
        "caption": "caption",
        "formula": "formula",
        "equation": "formula",
        "code": "code",
        "sheet": "sheet",
        "slide": "slide",
    }
    return mapping.get(normalized, "paragraph")


def _map_asset_kind(value: str) -> str:
    normalized = value.lower()
    return normalized if normalized in {"image", "chart", "page_render", "attachment"} else "image"


def configured_external_adapter(kind: ParserKind, worker_script: Path) -> ExternalWorkerDocumentAdapter:
    env_name = "E_REVIEW_DOCLING_PYTHON" if kind == "docling" else "E_REVIEW_MINERU_PYTHON"
    python_executable = os.getenv(env_name)
    if not python_executable:
        executable_name = "python.exe" if os.name == "nt" else "python"
        local_candidate = (
            Path(__file__).resolve().parents[2]
            / f".venv-{kind}"
            / ("Scripts" if os.name == "nt" else "bin")
            / executable_name
        )
        if local_candidate.is_file():
            python_executable = str(local_candidate)
    artifact_root = os.getenv("E_REVIEW_DOCUMENT_ARTIFACT_ROOT")
    timeout = int(os.getenv("E_REVIEW_DOCUMENT_PARSE_TIMEOUT_SECONDS", "600"))
    return ExternalWorkerDocumentAdapter(
        kind,
        python_executable=python_executable,
        worker_script=worker_script,
        artifact_root=artifact_root,
        timeout_seconds=timeout,
    )


def _external_worker_runtime_root(worker_script: Path) -> Path:
    configured = os.getenv("E_REVIEW_DOCUMENT_WORKER_ROOT", "").strip()
    root = Path(configured).expanduser() if configured else worker_script.resolve().parents[1] / "runtime" / "document-workers"
    if os.name == "nt" and not str(root.resolve()).isascii():
        root = Path(os.getenv("SystemDrive", "C:")) / "e-review-runtime" / "document-workers"
    root.mkdir(parents=True, exist_ok=True)
    return root.resolve()


def _external_worker_environment(kind: ParserKind, runtime_root: Path) -> dict[str, str]:
    environment = os.environ.copy()
    temp_root = runtime_root / f"{kind}-temp"
    home_root = runtime_root / f"{kind}-home"
    local_app_data = home_root / "AppData" / "Local"
    roaming_app_data = home_root / "AppData" / "Roaming"
    for path in (temp_root, local_app_data, roaming_app_data):
        path.mkdir(parents=True, exist_ok=True)

    environment.update(
        {
            "USERPROFILE": str(home_root),
            "HOME": str(home_root),
            "LOCALAPPDATA": str(local_app_data),
            "APPDATA": str(roaming_app_data),
            "TEMP": str(temp_root),
            "TMP": str(temp_root),
        }
    )
    if os.name == "nt":
        drive, tail = os.path.splitdrive(str(home_root))
        environment["HOMEDRIVE"] = drive
        environment["HOMEPATH"] = tail or "\\"
    if kind == "docling":
        environment["DOCLING_CACHE_DIR"] = str(home_root / ".cache" / "docling")
        environment["HF_HOME"] = str(home_root / ".cache" / "huggingface")
    return environment
