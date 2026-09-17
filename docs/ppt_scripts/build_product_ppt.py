from pathlib import Path
import datetime as dt
import re
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "ppt_assets"
SHOTS = ASSETS / "screenshots"
DIAGRAMS = ASSETS / "diagrams"
OUT = ROOT / "docs" / "ppt_output"
OUT.mkdir(parents=True, exist_ok=True)

PPTX = OUT / "E-Review-Agent-v104-Product-Intro.pptx"
PDF = OUT / "E-Review-Agent-v104-Product-Intro.pdf"
NOTES = OUT / "E-Review-Agent-v104-slide_notes.md"
INDEX = OUT / "E-Review-Agent-v104_screenshot_index.md"
MANIFEST = OUT / "E-Review-Agent-v104_asset_manifest.md"

W, H = 13.333, 7.5


def rel(p: Path) -> str:
    return str(p.relative_to(ROOT)).replace("\\", "/")


def register_pdf_font():
    for candidate in ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf"]:
        if Path(candidate).exists():
            pdfmetrics.registerFont(TTFont("DeckCJK", candidate))
            return "DeckCJK"
    return "Helvetica"


def add_textbox(slide, x, y, w, h, text, size=20, color="0F172A", bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def add_bullets(slide, x, y, w, h, bullets, size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for i, item in enumerate(bullets):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Microsoft YaHei"
        p.font.color.rgb = RGBColor.from_string("334155")
        p.space_after = Pt(7)
    return box


def add_panel(slide, x, y, w, h, title, body, fill="F8FAFC"):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.color.rgb = RGBColor.from_string("CBD5E1")
    add_textbox(slide, x + 0.18, y + 0.15, w - 0.36, 0.35, title, 16, "0F172A", True)
    add_textbox(slide, x + 0.18, y + 0.58, w - 0.36, h - 0.72, body, 12, "475569")


def add_image_fit(slide, image_path, x, y, w, h):
    image_path = Path(image_path)
    if not image_path.exists():
        add_panel(slide, x, y, w, h, "素材缺失", rel(image_path), "FEF2F2")
        return
    with Image.open(image_path) as img:
        iw, ih = img.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        width = w
        height = w / img_ratio
        left = x
        top = y + (h - height) / 2
    else:
        height = h
        width = h * img_ratio
        left = x + (w - width) / 2
        top = y
    slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def base_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor.from_string("F8FAFC")
    add_textbox(slide, 0.55, 0.32, 8.8, 0.48, title, 25, "0F172A", True)
    if subtitle:
        add_textbox(slide, 9.2, 0.42, 3.55, 0.35, subtitle, 11, "64748B", False, PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(0.92), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor.from_string("CBD5E1")
    line.line.color.rgb = RGBColor.from_string("CBD5E1")
    return slide


slides = [
    {
        "title": "E-Review Agent v1.0.4",
        "subtitle": "电商图文评论智能治理系统",
        "kind": "cover",
        "bullets": ["前后台一体化闭环", "Agentic RAG 产品能力", "可观测、可测试、可答辩展示"],
        "images": ["diagrams/architecture_overview.png"],
        "note": "本页开场介绍项目定位：E-Review Agent 是面向电商评论治理的毕业设计系统。v1.0.4 的重点不是单点 AI 分析，而是把真实用户评价、后台 Agent 巡检、风险任务和运营处理连成可演示、可验收的闭环。"
    },
    {
        "title": "业务背景与痛点",
        "bullets": ["电商评论包含文本、评分、图片链接和售后诉求，人工筛查容易遗漏。", "负向评论、图文不一致、退款售后等风险需要快速识别并闭环处理。", "传统后台只展示列表，缺少自动巡检、风险分级、处理反馈和过程可解释能力。"],
        "note": "这一页说明为什么要做这个系统。评论治理不是简单的关键词过滤，运营需要知道风险来自哪里、为什么被判高风险、后续由谁处理以及处理结果能否沉淀。"
    },
    {
        "title": "系统目标与产品定位",
        "bullets": ["定位：前后台一体化电商图文评论智能治理系统。", "目标一：真实 H5 用户端评价写入 litemall_comment。", "目标二：后台 Agent 自动巡检真实评价并生成风险任务。", "目标三：通过 Trace、Eval、Case Knowledge 支撑可解释与可验收。", "边界：不接真实支付、物流、退款，不声称生产级 SaaS。"],
        "note": "本页把系统边界讲清楚。项目聚焦评论治理与 Agent 闭环，不做真实支付、真实物流和真实退款，这样既符合毕设范围，也避免把答辩重点分散到基础商城能力上。"
    },
    {
        "title": "总体架构",
        "image": "diagrams/architecture_overview.png",
        "bullets": ["H5 用户端、管理后台、Spring Boot API、FastAPI AI 服务和 MySQL 共同组成。", "admin-api 与 wx-api 分离部署，AI 服务以 HTTP 能力接入。", "测试脚本覆盖服务状态、客户闭环、Agent 框架和最终验收。"],
        "note": "这页展示整体技术结构。前台负责产生真实业务数据，后台负责治理与运营，AI 服务提供分析能力，MySQL 保留分析结果、风险任务、运行轨迹和案例知识。"
    },
    {
        "title": "用户评价到 Agent 治理闭环",
        "image": "diagrams/customer_agent_loop.png",
        "bullets": ["浏览商品到发布评价形成真实业务入口。", "Agent 巡检扫描未分析的 litemall_comment。", "高风险评论进入风险中心，并在运营处理中心完成人工确认。"],
        "note": "这一页是核心演示路径。用户从前台产生评价，后台不需要手动复制文本，Agent 自动扫描真实评论并落库，运营人员在风险中心和处理中心完成闭环。"
    },
    {
        "title": "H5 用户端购买评价链路",
        "images": ["screenshots/01_h5_home.png", "screenshots/02_h5_demo_product_detail.png", "screenshots/03_h5_order_submit.png", "screenshots/07_h5_review_submit_success.png"],
        "bullets": ["展示首页、商品详情、订单提交和评价成功。", "演示支付用于状态推进，不调用真实支付网关。"],
        "note": "这里展示用户端入口。H5 页面可以浏览商品、进入详情、提交订单并发布图文评价。演示支付只是为了让订单进入可评价状态，不代表真实交易能力。"
    },
    {
        "title": "AI 工作台 Dashboard",
        "image": "screenshots/09_admin_ai_dashboard.png",
        "bullets": ["汇总待分析评论、风险任务、处理状态和趋势。", "用于答辩开场展示系统当前运行态。"],
        "note": "Dashboard 用来快速说明系统已经运行起来，能够统计分析结果、风险任务和处理状态。答辩时可以从这里进入各个子模块。"
    },
    {
        "title": "真实评价进入后台评论列表",
        "image": "screenshots/10_admin_comment_list_real_review.png",
        "bullets": ["H5 发布的评价进入原 litemall_comment。", "后台商品评论列表可以看到真实评价。", "一键 AI 分析仍可作为手动补充入口。"],
        "note": "这页证明系统不是只做后台模拟输入。用户端真实发布的商品评价会进入原商城评论表，后台商品评论列表可以查看，也可以手动触发 AI 分析。"
    },
    {
        "title": "Agent 巡检中心",
        "image": "screenshots/11_admin_patrol_center.png",
        "bullets": ["支持立即巡检和巡检日志查看。", "可扫描 demo_review 与 litemall_comment 两类来源。", "巡检结果写入 AI 分析表并触发风险任务。"],
        "note": "Agent 巡检中心是自动化治理入口。它把待分析评论批量扫描，调用 AI 服务完成分析，并记录巡检结果，便于答辩时展示自动化能力。"
    },
    {
        "title": "风险中心与运营处理中心",
        "images": ["screenshots/12_admin_risk_center.png", "screenshots/13_admin_operation_center.png"],
        "bullets": ["风险中心负责查看高风险评论和建议。", "运营处理中心负责人机协同决策、采纳建议和记录反馈。"],
        "note": "风险中心和运营处理中心构成治理闭环。AI 给出辅助建议，但最终处理由运营人员确认，这也体现了人机协同和系统边界。"
    },
    {
        "title": "Agentic RAG 工作流",
        "image": "diagrams/agentic_rag_workflow.png",
        "bullets": ["Review Analyst 解析评论与图片线索。", "Case Retriever 检索本地案例知识。", "Risk Decision 综合规则和案例证据输出建议。"],
        "note": "本页解释 v1.0.4 的 Agentic RAG 设计。系统保留本地案例检索和降级机制，不声称接入 Qdrant，而是强调可演示、可解释、可稳定验收的工程实现。"
    },
    {
        "title": "Agent Trace 可观测性",
        "images": ["screenshots/14_admin_agent_trace_list.png", "screenshots/15_admin_agent_trace_state_snapshot.png", "screenshots/16_admin_agent_trace_role_timeline.png"],
        "bullets": ["记录 Agent Run、步骤、状态快照和角色时间线。", "用于定位分析失败、解释决策链路和复盘运行过程。"],
        "note": "Agent Trace 解决了 AI 结果不可解释的问题。答辩时可以打开某一次运行，展示输入、步骤、状态和输出，证明系统不是黑盒结果。"
    },
    {
        "title": "Run Replay 与结果对比",
        "images": ["screenshots/17_admin_agent_trace_replay_compare.png", "diagrams/agent_trace_replay.png"],
        "bullets": ["对同一输入进行重放，比较原始运行和重放结果。", "用于演示 Agent 行为复盘和版本调整验证。"],
        "note": "Replay 能说明系统具备工程调试能力。即使是规则或 Mock AI，也可以通过重放机制观察流程变化，为后续接真实模型提供基础。"
    },
    {
        "title": "案例知识库与本地 RAG",
        "image": "screenshots/19_admin_case_knowledge_retrieval.png",
        "bullets": ["从历史风险和运营处理沉淀案例。", "分析时检索相似案例，补充处置依据。", "本阶段采用本地关键词检索与降级策略。"],
        "note": "案例知识库体现知识沉淀能力。系统把历史处理经验转化为可检索案例，用于辅助新评论判断，但不夸大为生产级向量数据库能力。"
    },
    {
        "title": "Agent Eval 质量评估",
        "image": "screenshots/18_admin_agent_eval_quality_health.png",
        "bullets": ["展示情感、风险类型、风险等级和案例检索指标。", "支持质量健康度和工具调用统计。", "用于支撑测试报告与答辩可信度。"],
        "note": "Agent Eval 用指标说明系统效果。它不是只说能分析，而是用样本集和通过率展示质量，让答辩老师能看到系统有测试依据。"
    },
    {
        "title": "诊断中心与配置状态",
        "images": ["screenshots/20_admin_diagnostics_failure_groups.png", "screenshots/21_admin_config_framework_status.png"],
        "bullets": ["诊断模块展示失败分组和健康状态。", "配置页展示当前 Agent 框架模式、RAG 状态和降级信息。"],
        "note": "这一页突出稳定性与可维护性。系统不仅有业务页面，还有诊断和配置状态页，便于演示前检查和出现问题时定位。"
    },
    {
        "title": "核心数据模型",
        "image": "diagrams/database_er_core.png",
        "bullets": ["litemall_comment 作为真实评价入口。", "analysis、risk_task、operation_log 形成治理主链路。", "agent_run、agent_step、case_knowledge 支撑可观测和 RAG。"],
        "note": "数据模型页用于论文和答辩技术说明。重点讲 source_type 和 source_id 如何区分真实评价、模拟评价和手动分析，避免重复分析并支撑来源追踪。"
    },
    {
        "title": "测试与验收结果",
        "images": ["diagrams/test_acceptance_matrix.png", "screenshots/22_final_acceptance_pass.png"],
        "bullets": ["最终验收脚本覆盖服务、接口、客户闭环、Agent 框架、质量和数据一致性。", "当前封版验收结论为 PASS，不伪造生产能力。"],
        "note": "本页展示验收证据。答辩时可以说明这些脚本是为了保证演示稳定，覆盖前后台服务、AI 服务、客户闭环和数据库一致性。"
    },
    {
        "title": "能力对比与创新点",
        "image": "diagrams/benchmark_matrix.png",
        "bullets": ["真实电商评价进入 Agent 自动治理，而不是孤立文本分析。", "Trace、Replay、Eval、Case Knowledge 提升可解释性。", "围绕毕业设计形成可演示、可测试、可写论文的完整系统。"],
        "note": "这一页总结创新点：真实业务闭环、自动巡检、风险运营闭环、运行可观测、本地 RAG 和质量评估。注意表述为毕业设计原型，不夸大商业化。"
    },
    {
        "title": "总结、边界与后续方向",
        "bullets": ["已完成：H5 真实评价、后台 Agent 巡检、风险任务、运营处理、Dashboard 联动。", "已完成：Trace、Replay、Eval、Case Knowledge、诊断、最终验收脚本。", "边界：未接真实支付、物流、退款；AI 服务仍可使用规则/Mock/降级模式。", "后续：接入真实模型、分布式任务锁、图片上传、向量数据库和更细粒度权限。"],
        "note": "最后收束项目价值和边界。强调当前版本适合毕业设计答辩和演示，后续可以继续接真实模型、真实图片上传、分布式巡检和更完整的生产级能力。"
    },
]


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    for idx, spec in enumerate(slides, start=1):
        if spec.get("kind") == "cover":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string("0F172A")
            add_textbox(slide, 0.75, 0.8, 7.5, 0.8, spec["title"], 40, "FFFFFF", True)
            add_textbox(slide, 0.78, 1.65, 7.0, 0.45, spec["subtitle"], 22, "CBD5E1")
            add_bullets(slide, 0.82, 2.5, 5.0, 2.0, spec["bullets"], 19)
            add_textbox(slide, 0.78, 6.7, 7.0, 0.3, "v1.0.4 Product Introduction / Graduation Defense", 12, "94A3B8")
            add_image_fit(slide, DIAGRAMS / "architecture_overview.png", 6.3, 1.35, 6.45, 4.8)
            continue
        slide = base_slide(prs, f"{idx:02d}. {spec['title']}", "E-Review Agent v1.0.4")
        if "image" in spec:
            add_image_fit(slide, ASSETS / spec["image"], 0.65, 1.25, 7.25, 5.55)
            add_bullets(slide, 8.25, 1.45, 4.35, 4.8, spec.get("bullets", []), 17)
        elif "images" in spec:
            imgs = spec["images"]
            if len(imgs) == 2:
                add_image_fit(slide, ASSETS / imgs[0], 0.7, 1.3, 5.7, 4.8)
                add_image_fit(slide, ASSETS / imgs[1], 6.65, 1.3, 5.7, 4.8)
                add_bullets(slide, 0.9, 6.25, 11.8, 0.75, spec.get("bullets", []), 13)
            elif len(imgs) == 3:
                add_image_fit(slide, ASSETS / imgs[0], 0.65, 1.25, 3.85, 4.85)
                add_image_fit(slide, ASSETS / imgs[1], 4.75, 1.25, 3.85, 4.85)
                add_image_fit(slide, ASSETS / imgs[2], 8.85, 1.25, 3.85, 4.85)
                add_bullets(slide, 0.9, 6.25, 11.8, 0.75, spec.get("bullets", []), 13)
            else:
                for i, img in enumerate(imgs[:4]):
                    x = 0.65 + (i % 4) * 3.05
                    add_image_fit(slide, ASSETS / img, x, 1.25, 2.8, 4.85)
                add_bullets(slide, 0.9, 6.25, 11.8, 0.75, spec.get("bullets", []), 13)
        else:
            add_bullets(slide, 1.0, 1.45, 11.2, 4.8, spec.get("bullets", []), 22)

    prs.save(PPTX)
    return prs


def build_notes():
    lines = ["# E-Review Agent v1.0.4 Product Intro Slide Notes", ""]
    for i, spec in enumerate(slides, start=1):
        lines.append(f"## Slide {i:02d}: {spec['title']}")
        lines.append(spec["note"])
        lines.append("")
    NOTES.write_text("\n".join(lines), encoding="utf-8")


def build_index_and_manifest():
    screenshot_names = [
        "01_h5_home.png", "02_h5_demo_product_detail.png", "03_h5_order_submit.png", "04_h5_demo_payment.png",
        "05_h5_demo_shipping.png", "06_h5_confirm_receipt.png", "07_h5_review_submit_success.png",
        "08_admin_login_or_dashboard.png", "09_admin_ai_dashboard.png", "10_admin_comment_list_real_review.png",
        "11_admin_patrol_center.png", "12_admin_risk_center.png", "13_admin_operation_center.png",
        "14_admin_agent_trace_list.png", "15_admin_agent_trace_state_snapshot.png", "16_admin_agent_trace_role_timeline.png",
        "17_admin_agent_trace_replay_compare.png", "18_admin_agent_eval_quality_health.png",
        "19_admin_case_knowledge_retrieval.png", "20_admin_diagnostics_failure_groups.png",
        "21_admin_config_framework_status.png", "22_final_acceptance_pass.png",
    ]
    index_lines = ["# E-Review Agent v1.0.4 Screenshot Index", "", "| File | Purpose |", "| --- | --- |"]
    for name in screenshot_names:
        index_lines.append(f"| docs/ppt_assets/screenshots/{name} | Product deck visual evidence |")
    INDEX.write_text("\n".join(index_lines) + "\n", encoding="utf-8")

    diagram_names = [
        "architecture_overview.png", "customer_agent_loop.png", "agentic_rag_workflow.png",
        "agent_trace_replay.png", "database_er_core.png", "test_acceptance_matrix.png", "benchmark_matrix.png",
    ]
    manifest = [
        "# E-Review Agent v1.0.4 Asset Manifest",
        "",
        f"Generated: {dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "## Screenshots",
    ]
    manifest.extend(f"- docs/ppt_assets/screenshots/{name}" for name in screenshot_names)
    manifest.extend(["", "## Diagrams"])
    manifest.extend(f"- docs/ppt_assets/diagrams/{name}" for name in diagram_names)
    manifest.extend(["", "## Outputs", f"- {rel(PPTX)}", f"- {rel(PDF)}", f"- {rel(NOTES)}", f"- {rel(INDEX)}"])
    manifest.extend(["", "## Generation Notes", "- Presentation generated with python-pptx because the local @oai/artifact-tool package was unavailable in the bundled runtime.", "- PDF generated directly with reportlab from the same slide plan."])
    MANIFEST.write_text("\n".join(manifest) + "\n", encoding="utf-8")


def build_pdf():
    font_name = register_pdf_font()
    c = canvas.Canvas(str(PDF), pagesize=landscape((960, 540)))
    width, height = landscape((960, 540))
    for i, spec in enumerate(slides, start=1):
        c.setFillColorRGB(0.97, 0.98, 0.99)
        c.rect(0, 0, width, height, fill=1, stroke=0)
        c.setFillColorRGB(0.06, 0.09, 0.16)
        c.setFont(font_name, 24)
        c.drawString(42, height - 52, f"{i:02d}. {spec['title']}")
        c.setStrokeColorRGB(0.8, 0.84, 0.9)
        c.line(42, height - 68, width - 42, height - 68)
        image_items = []
        if "image" in spec:
            image_items = [spec["image"]]
        elif "images" in spec:
            image_items = spec["images"][:2]
        if image_items:
            x = 48
            for img in image_items:
                p = ASSETS / img
                if p.exists():
                    c.drawImage(ImageReader(str(p)), x, 120, width=390, height=300, preserveAspectRatio=True, anchor="c")
                x += 430
        c.setFont(font_name, 14)
        c.setFillColorRGB(0.2, 0.25, 0.34)
        y = 102 if image_items else 390
        for bullet in spec.get("bullets", [])[:4]:
            c.drawString(60, y, f"- {bullet[:78]}")
            y -= 24
        c.setFont(font_name, 9)
        c.setFillColorRGB(0.45, 0.5, 0.58)
        c.drawRightString(width - 42, 24, "E-Review Agent v1.0.4")
        c.showPage()
    c.save()


def inspect_ppt_text():
    text = []
    with zipfile.ZipFile(PPTX) as zf:
        for name in zf.namelist():
            if name.startswith("ppt/slides/") and name.endswith(".xml"):
                xml = zf.read(name).decode("utf-8", errors="ignore")
                text.extend(re.findall(r"<a:t>(.*?)</a:t>", xml, flags=re.S))
    bad = ["http://", "https://", "C:\\", "D:\\", "锟", "�"]
    found = [item for item in bad if item in "\n".join(text)]
    if found:
        raise SystemExit(f"PPT_OUTPUT_CHECK_FAIL forbidden text: {found}")


if __name__ == "__main__":
    build_ppt()
    build_notes()
    build_pdf()
    build_index_and_manifest()
    inspect_ppt_text()
    print(f"PPT_OUTPUT_CHECK_PASS slides={len(slides)}")
    print(rel(PPTX))
    print(rel(PDF))
